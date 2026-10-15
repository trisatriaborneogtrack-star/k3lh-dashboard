"""
Dashboard Interaktif Tujuan, Sasaran & Program K3LH 2026
PT. Trisatria Persada Borneo

Sumber data : file "Program_Tujuan_Sasaran_K3LH_2026_Modern.xlsx" di Google Drive
              (sheet "Program Kerja 2026").
Tampilan    : identik dengan dashboard HTML (Tailwind + Chart.js), dirender di
              dalam Streamlit lewat components.html, datanya diisi dari Drive.
"""

import html
import io
import json
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl
import requests
import streamlit as st
import streamlit.components.v1 as components

# ---------------------------------------------------------------------------
# KONFIGURASI
# ---------------------------------------------------------------------------
# ID file di Google Drive (bagian setelah /d/ pada link).
# File harus di-share "Siapa saja yang memiliki link - Viewer".
FILE_ID = "1MxWz6wuBynEdhRlrGlH-z3-hCjCTk8DA"
SHEET_NAME = "Program Kerja 2026"   # nama tab sumber data
CACHE_SECONDS = 60                  # data ditarik ulang dari Drive tiap 60 detik
TZ_WITA = timezone(timedelta(hours=8))

MONTHS = ["JAN", "FEB", "MAR", "APR", "MEI", "JUN",
          "JUL", "AGU", "SEP", "OKT", "NOV", "DES"]

# File ini .xlsx yang disimpan di Drive (bukan Google Sheet asli), jadi
# export CSV (?format=csv&gid=...) selalu gagal 400. Yang berhasil: export xlsx.
DOWNLOAD_URLS = [
    f"https://docs.google.com/spreadsheets/d/{FILE_ID}/export?format=xlsx",
    f"https://drive.google.com/uc?export=download&id={FILE_ID}",
]

st.set_page_config(
    page_title="Dashboard Interaktif Tujuan, Sasaran & Program K3LH 2026",
    page_icon="🛡️",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# Hilangkan chrome Streamlit supaya halaman sama persis dengan versi HTML.
st.markdown(
    """
    <style>
      #MainMenu, header[data-testid="stHeader"], footer,
      [data-testid="stToolbar"], [data-testid="stDecoration"],
      [data-testid="stStatusWidget"] { display: none !important; }
      .stApp, [data-testid="stAppViewContainer"], .main { background: #0b0f19 !important; }
      .block-container, [data-testid="stMainBlockContainer"] {
          padding: 0 !important; max-width: 100% !important;
      }
      [data-testid="stVerticalBlock"] { gap: 0 !important; }
      iframe { display: block; }
    </style>
    """,
    unsafe_allow_html=True,
)


# ---------------------------------------------------------------------------
# AMBIL DATA
# ---------------------------------------------------------------------------
@st.cache_data(ttl=CACHE_SECONDS, show_spinner="Memuat data dari Google Drive…")
def fetch_workbook_bytes() -> tuple[bytes, str]:
    errors = []
    for url in DOWNLOAD_URLS:
        try:
            r = requests.get(url, timeout=30, allow_redirects=True)
            if r.status_code == 200 and r.content[:2] == b"PK":   # file xlsx = zip
                return r.content, datetime.now(TZ_WITA).strftime("%d %b %Y %H:%M WITA")
            errors.append(f"{url.split('?')[0]} -> HTTP {r.status_code}")
        except requests.RequestException as exc:
            errors.append(f"{url.split('?')[0]} -> {exc}")
    raise RuntimeError(
        "File tidak bisa diunduh. Pastikan file di Drive di-share "
        "'Siapa saja yang memiliki link - Viewer'. Detail: " + " | ".join(errors)
    )


# ---------------------------------------------------------------------------
# PARSING (berbasis judul kolom, bukan index tetap)
# ---------------------------------------------------------------------------
def _txt(v) -> str:
    return "" if v is None else re.sub(r"\s+", " ", str(v)).strip()


def _num(v) -> float:
    """Angka dari sel: dukung angka asli, '1.500.000', 'Rp 1,500,000', '-' dll."""
    if v is None or isinstance(v, bool):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s in ("", "-") or s.startswith("#"):
        return 0.0
    s = re.sub(r"[^0-9,.\-]", "", s)
    if not s:
        return 0.0
    if "," in s and "." in s:                     # 1.500.000,50 atau 1,500,000.50
        dec = "," if s.rfind(",") > s.rfind(".") else "."
        s = s.replace("." if dec == "," else ",", "").replace(",", ".")
    elif s.count(".") > 1 or re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")                    # format ribuan Indonesia
    elif s.count(",") > 1 or re.fullmatch(r"-?\d{1,3}(,\d{3})+", s):
        s = s.replace(",", "")
    else:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def status_from_ratio(ratio: float) -> str:
    # Sama dengan rumus kolom KETERCAPAIAN di Excel:
    # =IF(AA>=1,"TERPENUHI",IF(AA>0,"SEBAGIAN","BELUM"))
    if ratio >= 1:
        return "TERPENUHI"
    if ratio > 0:
        return "SEBAGIAN"
    return "BELUM"


def parse_workbook(content: bytes) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else None
    rows = None
    candidates = [ws] if ws is not None else [wb[n] for n in wb.sheetnames]
    for sheet in candidates:
        grid = [list(r) for r in sheet.iter_rows(values_only=True)]
        if any(_txt(c).upper() == "PROGRAM KERJA K3LH" for r in grid for c in r):
            rows = grid
            break
    if rows is None:
        raise RuntimeError(f"Tab berisi kolom 'PROGRAM KERJA K3LH' tidak ditemukan. Tab: {wb.sheetnames}")

    width = max(len(r) for r in rows)
    rows = [r + [None] * (width - len(r)) for r in rows]

    # 1) Baris judul utama
    h = next(i for i, r in enumerate(rows) if any(_txt(c).upper() == "PROGRAM KERJA K3LH" for c in r))
    head = [_txt(c).upper() for c in rows[h]]
    sub = [_txt(c).upper() for c in rows[h + 1]] if h + 1 < len(rows) else [""] * width

    def find_col(*keys):
        for i, t in enumerate(head):
            if any(k in t for k in keys):
                return i
        return None

    c_no = head.index("NO") if "NO" in head else 0
    c_prog = head.index("PROGRAM KERJA K3LH")
    c_ratio = find_col("TOTAL TARGET")
    c_status = find_col("KETERCAPAIAN")
    c_biaya = find_col("BIAYA")
    c_pic = find_col("PIC")

    # 2) Kolom Plan/Act tiap bulan dari sub-header "JAN Plan", "JAN Act", ...
    plan_cols, act_cols = [None] * 12, [None] * 12
    for i, t in enumerate(sub):
        for m, name in enumerate(MONTHS):
            if t.startswith(name):
                if "PLAN" in t:
                    plan_cols[m] = i
                elif "ACT" in t:
                    act_cols[m] = i
    if None in plan_cols or None in act_cols:     # fallback ke susunan standar C..Z
        plan_cols = [c_prog + 1 + 2 * m for m in range(12)]
        act_cols = [c_prog + 2 + 2 * m for m in range(12)]

    # 3) Baris program: kolom NO berisi angka & nama program terisi, sampai baris TOTAL
    programs, total_row = [], None
    for r in rows[h + 2:]:
        label = _txt(r[c_prog])
        if label.upper().startswith("TOTAL"):
            total_row = r
            break
        no_raw = r[c_no] if c_no is not None else None
        if not label or no_raw is None or _num(no_raw) <= 0:
            continue
        plans = [_num(r[c]) for c in plan_cols]
        acts = [_num(r[c]) for c in act_cols]
        tp, ta = sum(plans), sum(acts)
        ratio = ta / tp if tp > 0 else 0.0
        programs.append({
            "no": int(_num(no_raw)),
            "prog": label,
            "plans": plans,
            "acts": acts,
            "totPlan": tp,
            "totAct": ta,
            "ratio": ratio,
            "pct": round(ratio * 100),
            "status": status_from_ratio(ratio),
            "biaya": _num(r[c_biaya]) if c_biaya is not None else 0.0,
            "pic": _txt(r[c_pic]) if c_pic is not None else "",
            # nilai hasil rumus Excel, untuk validasi silang
            "_xl_ratio": r[c_ratio] if c_ratio is not None else None,
            "_xl_status": _txt(r[c_status]).upper() if c_status is not None else "",
        })

    if not programs:
        raise RuntimeError("Tidak ada baris program kerja yang terbaca.")

    m_plan = [sum(p["plans"][m] for p in programs) for m in range(12)]
    m_act = [sum(p["acts"][m] for p in programs) for m in range(12)]
    trend = [round(a / p * 100) if p > 0 else 0 for p, a in zip(m_plan, m_act)]
    tot_plan, tot_act = sum(m_plan), sum(m_act)
    counts = {s: sum(1 for p in programs if p["status"] == s) for s in ("TERPENUHI", "SEBAGIAN", "BELUM")}
    budget = sum(p["biaya"] for p in programs)
    avg_ratio = sum(p["ratio"] for p in programs) / len(programs)

    # 4) Validasi silang terhadap angka hasil rumus di file Excel
    issues = []
    for p in programs:
        xr = p["_xl_ratio"]
        if isinstance(xr, (int, float)) and abs(xr - p["ratio"]) > 1e-6:
            issues.append(f"No {p['no']}: ketercapaian {p['ratio']:.3f} vs Excel {xr:.3f}")
        if p["_xl_status"] and p["_xl_status"] != p["status"]:
            issues.append(f"No {p['no']}: status {p['status']} vs Excel {p['_xl_status']}")
    if total_row is not None:
        xl_plan = [_num(total_row[c]) for c in plan_cols]
        xl_act = [_num(total_row[c]) for c in act_cols]
        if [round(x) for x in xl_plan] != [round(x) for x in m_plan] or \
           [round(x) for x in xl_act] != [round(x) for x in m_act]:
            issues.append("Total per bulan berbeda dengan baris TOTAL di Excel")
        xl_avg = total_row[c_ratio] if c_ratio is not None else None
        if isinstance(xl_avg, (int, float)) and abs(xl_avg - avg_ratio) > 1e-6:
            issues.append(f"Rata-rata ketercapaian {avg_ratio:.3f} vs Excel {xl_avg:.3f}")
        if c_biaya is not None and total_row[c_biaya] is not None and \
           abs(_num(total_row[c_biaya]) - budget) > 0.5:
            issues.append(f"Total biaya {budget:,.0f} vs Excel {_num(total_row[c_biaya]):,.0f}")

    for p in programs:
        p.pop("_xl_ratio"); p.pop("_xl_status")

    return {
        "programs": programs,
        "mPlan": m_plan,
        "mAct": m_act,
        "mTrend": trend,
        "totalPlan": tot_plan,
        "totalAct": tot_act,
        # = sel TOTAL TARGET pada baris TOTAL di Excel: rata-rata ketercapaian 26 program
        "overallPct": round(avg_ratio * 100),
        "counts": counts,
        "budget": budget,
        "pics": sorted({p["pic"] for p in programs if p["pic"]}),
        "validation": {"ok": not issues, "issues": issues, "n": len(programs)},
    }


# ---------------------------------------------------------------------------
# RENDER
# ---------------------------------------------------------------------------
def estimate_height(n: int) -> int:
    # header+KPI ~330, baris chart ~420, tren ~360, tabel summary & matriks
    return 330 + 420 + 360 + (170 + n * 45) + (230 + n * 41 + 100) + 80


try:
    content, synced_at = fetch_workbook_bytes()
    data = parse_workbook(content)
    data["syncedAt"] = synced_at

    template = (Path(__file__).parent / "dashboard_template.html").read_text(encoding="utf-8")
    payload = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
    page = template.replace("__DASHBOARD_DATA__", payload)

    if hasattr(st, "iframe"):          # Streamlit baru: tinggi iframe mengikuti isi
        st.iframe(page, height="content")
    else:                              # Streamlit lama
        components.html(page, height=estimate_height(len(data["programs"])), scrolling=False)

except Exception as exc:  # tampilkan error yang jelas, bukan halaman kosong
    st.markdown(
        f"""<div style="margin:24px;padding:20px;border-radius:16px;background:#111827;
        border:1px solid #7f1d1d;color:#fecaca;font-family:Inter,sans-serif">
        <b>Gagal memuat data dashboard.</b><br><span style="font-size:13px">{html.escape(str(exc))}</span></div>""",
        unsafe_allow_html=True,
    )
