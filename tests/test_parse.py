"""Uji parser terhadap salinan isi file Excel asli (dibaca 15 Okt 2026).

Jalankan:  python tests/test_parse.py
"""
import io
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

# Nilai mentah dari sheet "Program Kerja 2026": NO, PROGRAM, 24 sel plan/act, BIAYA, PIC,
# lalu nilai hasil rumus Excel kolom AA (TOTAL TARGET) & AB (KETERCAPAIAN).
ROWS = [
    (1, "Pembuatan Materi Safety Training", "1 1 1 1 1 7 1 5 1 0 1 0 1 0 1 0 1 0 1 0 1 0 1 0", 0, "Leo ", 1.1666666666666667, "TERPENUHI"),
    (2, "Conduct Safety Training", "1 1 1 0 1 1 1 0 1 0 1 0 1 0 1 0 1 0 1 0 1 0 1 0", 0, "Leo", 0.16666666666666666, "SEBAGIAN"),
    (3, "Cunduct K3 Induction", "1 0 1 0 1 1 1 2 1 1 1 0 1 1 1 1 1 0 1 0 1 0 1 0", 0, "Leo ", 0.5, "SEBAGIAN"),
    (4, "Pembuatan IBPR", "0 0 0 0 0 0 0 0 0 0 1 0 0 0 0 0 0 0 0 0 0 0 0 0", 0, "Leo ", 0, "BELUM"),
    (5, "Observasi Terencana Karyawan", "0 0 0 0 1 0 0 1 0 0 1 0 0 0 0 0 1 0 0 0 0 0 1 0", 0, "Leo", 0.25, "SEBAGIAN"),
    (6, "Pembentukan Team Tanggap Darurat", "0 0 0 0 0 0 1 0 0 0 0 0 0 0 0 0 0 1 0 0 0 0 0 0", 0, "Leo", 1, "TERPENUHI"),
    (7, "Simulasi Tanggap Darurat", "0 0 0 0 0 0 0 0 0 0 1 0 0 0 0 0 0 0 0 0 0 0 0 0", 500000, "Leo", 0, "BELUM"),
    (8, "Pengadaan APD (HO)", "0 0 0 0 1 1 0 0 0 1 1 0 0 0 0 1 1 0 0 0 0 0 0 0", 8750000, "Leo", 1, "TERPENUHI"),
    (9, "Refill APAR", "0 0 0 0 0 0 0 0 0 0 0 0 1 0 0 0 0 0 0 0 0 0 0 0", 2000000, "Leo", 0, "BELUM"),
    (10, "Report Stok APD", "1 0 1 1 1 1 1 1 1 1 1 1 1 1 1 1 1 1 1 1 1 0 1 0", 0, None, 0.75, "SEBAGIAN"),
    (11, "Inspeksi APD (Personel)", "1 0 1 0 1 4 1 0 1 0 1 0 1 0 1 0 1 0 1 0 1 0 1 0", 0, None, 0.3333333333333333, "SEBAGIAN"),
    (12, "Inspeksi APAR", "1 0 0 0 0 1 1 1 0 0 0 0 1 1 0 0 0 0 1 0 0 0 0 0", 0, None, 0.75, "SEBAGIAN"),
    (13, "Inspeksi P3K", "1 0 0 0 0 1 1 1 0 0 0 0 1 1 0 0 0 0 1 0 0 0 0 0", 0, None, 0.75, "SEBAGIAN"),
    (14, "Inspeksi Genset", "1 0 1 0 1 1 1 1 1 0 1 0 1 1 1 0 1 0 1 0 1 0 1 0", 0, None, 0.25, "SEBAGIAN"),
    (15, "Safety Talk", "2 0 2 2 2 5 2 6 2 1 2 2 2 2 2 2 2 2 2 0 2 0 2 0", 0, "Leo ", 0.9166666666666666, "SEBAGIAN"),
    (16, "Poster K3", "2 0 2 1 2 3 2 3 2 3 2 4 2 4 2 4 2 3 2 1 2 0 2 0", 200000, "Leo ", 1.0833333333333333, "TERPENUHI"),
    (17, "Evaluasi Dokumen & Prosedur", "0 0 0 0 0 0 1 0 0 0 0 0 0 0 1 0 0 0 0 0 0 0 1 0", 0, "Leo ", 0, "BELUM"),
    (18, "Laporan Pengecekan Kendaraan Operasional (P2H)", "1 0 1 0 1 0 1 1 1 0 1 0 1 0 1 0 1 0 1 0 1 0 1 0", 0, "Leo ", 0.08333333333333333, "SEBAGIAN"),
    (19, "Literasi SOP", "0 0 0 0 0 0 0 0 1 0 0 0 0 0 0 0 1 0 0 0 0 0 0 0", 0, "Leo", 0, "BELUM"),
    (20, "Evaluasi Perundangan secara berkala disesuaikan dengan JDIH atau sumber lainnya", "0 0 0 0 1 1 0 0 0 0 1 0 0 0 0 0 1 0 0 0 0 0 1 0", 0, "Leo ", 0.25, "SEBAGIAN"),
    (21, "House keeping di Area Kerja", "1 0 1 0 1 1 1 0 1 0 1 0 1 0 1 0 1 0 1 0 1 0 1 0", 0, "Leo ", 0.08333333333333333, "SEBAGIAN"),
    (22, "Rambu - rambu K3", "0 0 0 0 0 0 0 1 0 0 0 0 0 0 1 0 0 0 0 0 0 0 0 0", 500000, "Leo ", 1, "TERPENUHI"),
    (23, "Program Olahraga Karyawan", "1 1 1 2 1 0 1 1 1 1 1 1 1 0 1 1 1 0 1 0 1 0 1 0", 6000000, "Leo", 0.5833333333333334, "SEBAGIAN"),
    (24, "Management Visit", "0 0 0 0 0 0 0 1 0 0 0 0 1 0 0 0 0 0 0 0 0 0 1 0", 4000000, "Leo", 0.5, "SEBAGIAN"),
    (25, "Pembuatan RKAB Project Support", "0 0 0 0 1 1 0 0 0 0 0 0 0 0 0 0 0 0 0 0 0 0 1 0", 75730000, "Leo", 0.5, "SEBAGIAN"),
    (26, "Training APAR", "0 0 0 0 1 0 0 0 0 0 0 0 0 0 0 0 0 0 0 0 0 0 0 0", 500000, "Leo", 0, "BELUM"),
]
TOTAL_ROW = "15 3 13 7 18 29 17 25 14 8 18 8 17 11 15 10 17 7 15 2 13 0 18 0"
MONTHS = ["JAN", "FEB", "MAR", "APR", "MEI", "JUN", "JUL", "AGU", "SEP", "OKT", "NOV", "DES"]


def build_fixture() -> bytes:
    wb = openpyxl.Workbook()
    wb.active.title = "Dashboard Pencapaian"
    ws = wb.create_sheet("Program Kerja 2026")
    ws["A1"] = "PT. TRISATRIA PERSADA BORNEO"
    ws["A2"] = "PROGRAM KERJA, TUJUAN DAN SASARAN K3LH - PERIODE TAHUN 2026"
    head = {"A3": "NO", "B3": "PROGRAM KERJA K3LH", "C3": "QUARTER I", "I3": "QUARTER II", "O3": "QUARTER III",
            "U3": "QUARTER IV", "AA3": "TOTAL TARGET", "AB3": "KETERCAPAIAN", "AC3": "RINCIAN BIAYA (Rp)", "AD3": "PIC"}
    for k, v in head.items():
        ws[k] = v
    for m, name in enumerate(MONTHS):
        ws.cell(4, 3 + 2 * m, f"{name}\nPlan")
        ws.cell(4, 4 + 2 * m, f"{name}\nAct")
    for i, (no, prog, vals, biaya, pic, xr, xs) in enumerate(ROWS):
        r = 5 + i
        ws.cell(r, 1, no); ws.cell(r, 2, prog)
        for j, v in enumerate(vals.split()):
            ws.cell(r, 3 + j, int(v))
        ws.cell(r, 27, xr); ws.cell(r, 28, xs); ws.cell(r, 29, biaya); ws.cell(r, 30, pic)
    ws.cell(31, 2, "TOTAL TARGET PER BULAN (PLAN / ACTUAL)")
    for j, v in enumerate(TOTAL_ROW.split()):
        ws.cell(31, 3 + j, int(v))
    ws.cell(31, 27, 0.4583333333333333); ws.cell(31, 29, 98180000)
    ws.cell(32, 2, "PERSENTASE KETERCAPAIAN PER BULAN")
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def main():
    import app  # noqa: E402  (Streamlit "bare mode" saat diimpor tanpa server)

    d = app.parse_workbook(build_fixture())
    assert d["validation"]["ok"], d["validation"]["issues"]
    assert len(d["programs"]) == 26
    assert d["counts"] == {"TERPENUHI": 5, "SEBAGIAN": 15, "BELUM": 6}, d["counts"]
    assert d["budget"] == 98_180_000
    assert d["mPlan"] == [15, 13, 18, 17, 14, 18, 17, 15, 17, 15, 13, 18]
    assert d["mAct"] == [3, 7, 29, 25, 8, 8, 11, 10, 7, 2, 0, 0]
    assert d["overallPct"] == 46
    assert d["mTrend"] == [20, 54, 161, 147, 57, 44, 65, 67, 41, 13, 0, 0], d["mTrend"]
    assert d["programs"][16]["prog"] == "Evaluasi Dokumen & Prosedur"
    assert d["programs"][9]["pic"] == ""
    print("OK — 26 program,", d["counts"], "anggaran", f"{d['budget']:,.0f}",
          "| rata-rata", f"{d['overallPct']}%", "| tren", d["mTrend"])


if __name__ == "__main__":
    main()
